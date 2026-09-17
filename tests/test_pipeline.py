from unittest.mock import MagicMock

from lead_collector.discovery.models import DiscoveryResult
from lead_collector.extraction.exceptions import WebsiteFetchError
from lead_collector.extraction.website import FetchResult
from lead_collector.models import ValidationStatus
from lead_collector.pipeline import LeadCollectionPipeline
from lead_collector.storage.database import LeadDatabase
from lead_collector.storage.repository import LeadRepository
from lead_collector.export.csv_exporter import CSVLeadExporter
from lead_collector.export.excel_exporter import ExcelLeadExporter


def make_result(
    url: str,
    position: int = 1,
) -> DiscoveryResult:
    return DiscoveryResult(
        title="Example Company",
        url=url,
        snippet="Example company website",
        source="test",
        position=position,
    )


def test_pipeline_processes_discovered_results():
    discovery = MagicMock()

    discovery.search.return_value = [
        make_result("https://example.com"),
    ]

    fetcher = MagicMock()

    fetcher.fetch.return_value = FetchResult(
        url="https://example.com",
        status_code=200,
        content="""
            <html>
                <head>
                    <title>Example Company</title>
                </head>
                <body>
                    Contact sales@example.com
                </body>
            </html>
        """,
    )

    pipeline = LeadCollectionPipeline(
        discovery_provider=discovery,
        website_fetcher=fetcher,
    )

    result = pipeline.run(
        "software companies",
        max_results=10,
    )

    assert result.discovered == 1
    assert result.unique_results == 1
    assert result.fetched == 1
    assert result.failed_fetches == 0
    assert len(result.leads) == 1

    lead = result.leads[0]

    assert lead.company_name == "Example Company"
    assert str(lead.website) == "https://example.com/"
    assert str(lead.email) == "sales@example.com"
    assert lead.validation_status == ValidationStatus.VALID
    assert lead.lead_score == 65


def test_pipeline_deduplicates_results():
    discovery = MagicMock()

    discovery.search.return_value = [
        make_result("https://example.com"),
        make_result("https://example.com/"),
        make_result("https://other.com"),
    ]

    fetcher = MagicMock()

    fetcher.fetch.side_effect = [
        FetchResult(
            url="https://example.com",
            status_code=200,
            content="<title>Example</title>",
        ),
        FetchResult(
            url="https://other.com",
            status_code=200,
            content="<title>Other</title>",
        ),
    ]

    pipeline = LeadCollectionPipeline(
        discovery_provider=discovery,
        website_fetcher=fetcher,
    )

    result = pipeline.run("companies")

    assert result.discovered == 3
    assert result.unique_results == 2
    assert result.fetched == 2
    assert len(result.leads) == 2

    assert fetcher.fetch.call_count == 2


def test_pipeline_continues_after_fetch_failure():
    discovery = MagicMock()

    discovery.search.return_value = [
        make_result("https://broken.example"),
        make_result("https://working.example"),
    ]

    fetcher = MagicMock()

    fetcher.fetch.side_effect = [
        WebsiteFetchError("Website fetch failed."),
        FetchResult(
            url="https://working.example",
            status_code=200,
            content="<title>Working Company</title>",
        ),
    ]

    pipeline = LeadCollectionPipeline(
        discovery_provider=discovery,
        website_fetcher=fetcher,
    )

    result = pipeline.run("companies")

    assert result.discovered == 2
    assert result.unique_results == 2
    assert result.fetched == 1
    assert result.failed_fetches == 1
    assert len(result.leads) == 1
    assert result.leads[0].company_name == "Working Company"


def test_pipeline_handles_empty_discovery_results():
    discovery = MagicMock()
    discovery.search.return_value = []

    fetcher = MagicMock()

    pipeline = LeadCollectionPipeline(
        discovery_provider=discovery,
        website_fetcher=fetcher,
    )

    result = pipeline.run("companies")

    assert result.discovered == 0
    assert result.unique_results == 0
    assert result.fetched == 0
    assert result.failed_fetches == 0
    assert result.leads == []

    fetcher.fetch.assert_not_called()


def test_pipeline_persists_leads_to_database(tmp_path):
    discovery = MagicMock()
    discovery.search.return_value = [
        make_result("https://example.com"),
    ]

    fetcher = MagicMock()
    fetcher.fetch.return_value = FetchResult(
        url="https://example.com",
        status_code=200,
        content="""
            <html>
                <head>
                    <title>Example Company</title>
                </head>
                <body>
                    Contact sales@example.com
                </body>
            </html>
        """,
    )

    database = LeadDatabase(tmp_path / "leads.db")

    pipeline = LeadCollectionPipeline(
        discovery_provider=discovery,
        website_fetcher=fetcher,
        database=database,
    )

    result = pipeline.run("software companies")

    assert len(result.leads) == 1
    assert database.count() == 1

    stored_leads = LeadRepository(
        tmp_path / "leads.db"
    ).get_all()

    assert len(stored_leads) == 1
    assert stored_leads[0].company_name == "Example Company"
    assert str(stored_leads[0].email) == "sales@example.com"


def test_pipeline_does_not_persist_failed_fetches(tmp_path):
    discovery = MagicMock()
    discovery.search.return_value = [
        make_result("https://broken.example"),
    ]

    fetcher = MagicMock()
    fetcher.fetch.side_effect = WebsiteFetchError(
        "Website fetch failed."
    )

    database = LeadDatabase(tmp_path / "leads.db")

    pipeline = LeadCollectionPipeline(
        discovery_provider=discovery,
        website_fetcher=fetcher,
        database=database,
    )

    result = pipeline.run("companies")

    assert result.failed_fetches == 1
    assert result.leads == []
    assert database.count() == 0


def test_pipeline_persists_multiple_leads(tmp_path):
    discovery = MagicMock()
    discovery.search.return_value = [
        make_result("https://company-one.example", position=1),
        make_result("https://company-two.example", position=2),
    ]

    fetcher = MagicMock()
    fetcher.fetch.side_effect = [
        FetchResult(
            url="https://company-one.example",
            status_code=200,
            content="""
                <html>
                    <head>
                        <title>Company One</title>
                    </head>
                    <body>
                        Contact one@companyone.com
                    </body>
                </html>
            """,
        ),
        FetchResult(
            url="https://company-two.example",
            status_code=200,
            content="""
                <html>
                    <head>
                        <title>Company Two</title>
                    </head>
                    <body>
                        Contact two@companytwo.com
                    </body>
                </html>
            """,
        ),
    ]

    database = LeadDatabase(tmp_path / "leads.db")

    pipeline = LeadCollectionPipeline(
        discovery_provider=discovery,
        website_fetcher=fetcher,
        database=database,
    )

    result = pipeline.run("software companies")

    assert result.discovered == 2
    assert result.unique_results == 2
    assert result.fetched == 2
    assert result.failed_fetches == 0
    assert len(result.leads) == 2
    assert database.count() == 2

    stored_leads = LeadRepository(
        tmp_path / "leads.db"
    ).get_all()

    assert [lead.company_name for lead in stored_leads] == [
        "Company One",
        "Company Two",
    ]


def test_pipeline_exports_csv_and_excel(tmp_path):
    discovery = MagicMock()
    discovery.search.return_value = [
        make_result("https://example.com"),
    ]

    fetcher = MagicMock()
    fetcher.fetch.return_value = FetchResult(
        url="https://example.com",
        status_code=200,
        content="""
            <html>
                <head>
                    <title>Example Company</title>
                </head>
                <body>
                    Contact sales@example.com
                </body>
            </html>
        """,
    )

    csv_exporter = CSVLeadExporter()
    excel_exporter = ExcelLeadExporter()

    csv_path = tmp_path / "exports" / "leads.csv"
    excel_path = tmp_path / "exports" / "leads.xlsx"

    pipeline = LeadCollectionPipeline(
        discovery_provider=discovery,
        website_fetcher=fetcher,
        csv_exporter=csv_exporter,
        excel_exporter=excel_exporter,
    )

    result = pipeline.run(
        "software companies",
        csv_path=csv_path,
        excel_path=excel_path,
    )

    assert len(result.leads) == 1
    assert csv_path.exists()
    assert excel_path.exists()


def test_pipeline_exports_lead_data(tmp_path):
    discovery = MagicMock()
    discovery.search.return_value = [
        make_result("https://example.com"),
    ]

    fetcher = MagicMock()
    fetcher.fetch.return_value = FetchResult(
        url="https://example.com",
        status_code=200,
        content="""
            <html>
                <head>
                    <title>Example Company</title>
                </head>
                <body>
                    Contact sales@example.com
                </body>
            </html>
        """,
    )

    csv_exporter = CSVLeadExporter()
    excel_exporter = ExcelLeadExporter()

    csv_path = tmp_path / "leads.csv"
    excel_path = tmp_path / "leads.xlsx"

    pipeline = LeadCollectionPipeline(
        discovery_provider=discovery,
        website_fetcher=fetcher,
        csv_exporter=csv_exporter,
        excel_exporter=excel_exporter,
    )

    result = pipeline.run(
        "software companies",
        csv_path=csv_path,
        excel_path=excel_path,
    )

    assert len(result.leads) == 1

    # Verify CSV content.
    csv_content = csv_path.read_text(encoding="utf-8")

    assert "Example Company" in csv_content
    assert "sales@example.com" in csv_content

    # Verify Excel content.
    from openpyxl import load_workbook

    workbook = load_workbook(excel_path)
    worksheet = workbook["Leads"]

    headers = [cell.value for cell in worksheet[1]]
    company_index = headers.index("Company Name") + 1
    email_index = headers.index("Email") + 1

    assert worksheet.cell(row=2, column=company_index).value == "Example Company"
    assert worksheet.cell(row=2, column=email_index).value == "sales@example.com"


def test_pipeline_skips_challenge_pages():
    discovery = MagicMock()
    discovery.search.return_value = [
        make_result("https://example.com"),
    ]

    fetcher = MagicMock()
    fetcher.fetch.return_value = FetchResult(
        url="https://example.com",
        status_code=200,
        content="""
            <html>
                <head>
                    <title>Radware Captcha Page</title>
                </head>
                <body>
                    Please complete the CAPTCHA to continue.
                </body>
            </html>
        """,
        is_challenge_page=True,
    )

    parser = MagicMock()
    extractor = MagicMock()
    validator = MagicMock()

    pipeline = LeadCollectionPipeline(
        discovery_provider=discovery,
        website_fetcher=fetcher,
        html_parser=parser,
        lead_extractor=extractor,
        lead_validator=validator,
    )

    result = pipeline.run("software companies")

    assert result.discovered == 1
    assert result.unique_results == 1
    assert result.fetched == 1
    assert result.failed_fetches == 0
    assert result.leads == []

    parser.parse.assert_not_called()
    extractor.extract.assert_not_called()
    validator.validate.assert_not_called()


def test_pipeline_processes_normal_page_when_not_challenge():
    discovery = MagicMock()
    discovery.search.return_value = [
        make_result("https://example.com"),
    ]

    fetcher = MagicMock()
    fetcher.fetch.return_value = FetchResult(
        url="https://example.com",
        status_code=200,
        content="<title>Example Company</title>",
        is_challenge_page=False,
    )

    pipeline = LeadCollectionPipeline(
        discovery_provider=discovery,
        website_fetcher=fetcher,
    )

    result = pipeline.run("software companies")

    assert result.fetched == 1
    assert result.failed_fetches == 0
    assert len(result.leads) == 1
    assert result.leads[0].company_name == "Example Company"