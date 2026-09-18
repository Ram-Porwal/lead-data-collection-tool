from openpyxl import load_workbook

from lead_collector.export.excel_exporter import ExcelLeadExporter
from lead_collector.models import Lead, ValidationStatus


def create_lead() -> Lead:
    return Lead(
        company_name="Example Technologies",
        website="https://example.com",
        industry="Software",
        city="Ahmedabad",
        state="Gujarat",
        country="India",
        contact_name="John Doe",
        contact_role="Sales Manager",
        email="sales@example.com",
        phone="+91 98765 43210",
        phone_country="IN",
        linkedin_url=(
            "https://www.linkedin.com/company/example"
        ),
        source_url="https://example.com",
        lead_score=95,
        validation_status=ValidationStatus.VALID,
    )


def test_export_creates_excel_file(tmp_path):
    output_path = tmp_path / "exports" / "leads.xlsx"

    result = ExcelLeadExporter().export(
        [create_lead()],
        output_path,
    )

    assert result == output_path
    assert output_path.exists()


def test_export_writes_lead_data(tmp_path):
    output_path = tmp_path / "leads.xlsx"

    ExcelLeadExporter().export(
        [create_lead()],
        output_path,
    )

    workbook = load_workbook(output_path)
    worksheet = workbook["Leads"]

    assert worksheet["A2"].value == "Example Technologies"
    assert worksheet["C2"].value == "Software"
    assert worksheet["I2"].value == "sales@example.com"
    assert worksheet["J2"].value == "+91 98765 43210"
    assert worksheet["K2"].value == "IN"
    assert worksheet["N2"].value == "valid"
    assert worksheet["M2"].value == 95
    assert worksheet["N2"].value is not None


def test_export_formats_headers(tmp_path):
    output_path = tmp_path / "leads.xlsx"

    ExcelLeadExporter().export(
        [create_lead()],
        output_path,
    )

    workbook = load_workbook(output_path)
    worksheet = workbook["Leads"]

    assert worksheet["A1"].value == "Company Name"
    assert worksheet["B1"].value == "Website"
    assert worksheet["I1"].value == "Email"
    assert worksheet["M1"].value == "Lead Score"
    assert worksheet["K1"].value == "Phone Country"
    assert worksheet["P1"].value == "Created At"
    assert worksheet["Q1"].value == "ID"
    assert worksheet["A1"].font.bold is True
    assert worksheet["A1"].font.color.rgb == "00FFFFFF"


def test_export_freezes_header(tmp_path):
    output_path = tmp_path / "leads.xlsx"

    ExcelLeadExporter().export(
        [create_lead()],
        output_path,
    )

    workbook = load_workbook(output_path)
    worksheet = workbook["Leads"]

    assert worksheet.freeze_panes == "A2"


def test_export_creates_excel_table(tmp_path):
    output_path = tmp_path / "leads.xlsx"

    ExcelLeadExporter().export(
        [create_lead()],
        output_path,
    )

    workbook = load_workbook(output_path)
    worksheet = workbook["Leads"]

    assert "LeadReport" in worksheet.tables
    assert worksheet.tables["LeadReport"].ref == "A1:Q2"


def test_export_creates_hyperlinks(tmp_path):
    output_path = tmp_path / "leads.xlsx"

    ExcelLeadExporter().export(
        [create_lead()],
        output_path,
    )

    workbook = load_workbook(output_path)
    worksheet = workbook["Leads"]

    assert worksheet["B2"].value == "Open Website"
    assert worksheet["B2"].hyperlink.target == "https://example.com/"

    assert worksheet["L2"].value == "LinkedIn"
    assert (
        worksheet["L2"].hyperlink.target
        == "https://www.linkedin.com/company/example"
    )

    assert worksheet["O2"].value == "View Source"
    assert worksheet["O2"].hyperlink.target == "https://example.com/"


def test_export_formats_created_at(tmp_path):
    output_path = tmp_path / "leads.xlsx"

    ExcelLeadExporter().export(
        [create_lead()],
        output_path,
    )

    workbook = load_workbook(output_path)
    worksheet = workbook["Leads"]

    assert worksheet["P2"].number_format == "yyyy-mm-dd hh:mm"


def test_export_handles_empty_leads(tmp_path):
    output_path = tmp_path / "leads.xlsx"

    ExcelLeadExporter().export([], output_path)

    workbook = load_workbook(output_path)
    worksheet = workbook["Leads"]

    assert worksheet.max_row == 1
    assert worksheet["A1"].value == "Company Name"
    assert worksheet["Q1"].value == "ID"