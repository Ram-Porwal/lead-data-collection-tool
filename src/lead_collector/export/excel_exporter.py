from datetime import timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from lead_collector.models import Lead


class ExcelLeadExporter:
    """Export leads to a client-friendly formatted Excel workbook."""

    FIELDNAMES = [
        "company_name",
        "website",
        "industry",
        "city",
        "state",
        "country",
        "contact_name",
        "contact_role",
        "email",
        "phone",
        "linkedin_url",
        "lead_score",
        "validation_status",
        "source_url",
        "created_at",
        "id",
    ]

    DISPLAY_NAMES = {
        "company_name": "Company Name",
        "website": "Website",
        "industry": "Industry",
        "city": "City",
        "state": "State",
        "country": "Country",
        "contact_name": "Contact Name",
        "contact_role": "Contact Role",
        "email": "Email",
        "phone": "Phone",
        "linkedin_url": "LinkedIn",
        "lead_score": "Lead Score",
        "validation_status": "Validation Status",
        "source_url": "Source",
        "created_at": "Created At",
        "id": "ID",
    }

    LINK_COLUMNS = {
        "website",
        "linkedin_url",
        "source_url",
    }

    WIDTHS = {
        "company_name": 28,
        "website": 18,
        "industry": 20,
        "city": 16,
        "state": 16,
        "country": 16,
        "contact_name": 22,
        "contact_role": 20,
        "email": 30,
        "phone": 20,
        "linkedin_url": 18,
        "lead_score": 12,
        "validation_status": 20,
        "source_url": 18,
        "created_at": 22,
        "id": 38,
    }

    def export(
        self,
        leads: list[Lead],
        output_path: str | Path,
    ) -> Path:
        """Export leads to a formatted Excel workbook."""

        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Leads"

        self._write_headers(worksheet)
        self._write_leads(worksheet, leads)
        self._format_worksheet(worksheet)

        workbook.save(path)

        return path

    def _write_headers(self, worksheet) -> None:
        """Write and format the worksheet headers."""

        headers = [
            self.DISPLAY_NAMES[field]
            for field in self.FIELDNAMES
        ]

        for column, header in enumerate(headers, start=1):
            worksheet.cell(
                row=1,
                column=column,
                value=header,
            )

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )
        header_font = Font(
            bold=True,
            color="FFFFFF",
        )

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        worksheet.row_dimensions[1].height = 24

    def _write_leads(
        self,
        worksheet,
        leads: list[Lead],
    ) -> None:
        """Write lead records to the worksheet."""

        for lead in leads:
            values = {
                "company_name": lead.company_name,
                "website": str(lead.website) if lead.website else "",
                "industry": lead.industry or "",
                "city": lead.city or "",
                "state": lead.state or "",
                "country": lead.country or "",
                "contact_name": lead.contact_name or "",
                "contact_role": lead.contact_role or "",
                "email": str(lead.email) if lead.email else "",
                "phone": lead.phone or "",
                "linkedin_url": (
                    str(lead.linkedin_url)
                    if lead.linkedin_url
                    else ""
                ),
                "lead_score": lead.lead_score,
                "validation_status": lead.validation_status.value,
                "source_url": (
                    str(lead.source_url)
                    if lead.source_url
                    else ""
                ),
                "created_at": (
                    lead.created_at.astimezone(timezone.utc).replace(tzinfo=None)
                    if lead.created_at
                    else None
                ),
                "id": str(lead.id),
            }

            worksheet.append(
                [
                    values[field]
                    for field in self.FIELDNAMES
                ]
            )

        for row in worksheet.iter_rows(
            min_row=2,
            max_row=worksheet.max_row,
        ):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    def _format_worksheet(self, worksheet) -> None:
        """Apply usability and presentation formatting."""

        worksheet.freeze_panes = "A2"

        if worksheet.max_row >= 2:
            table_ref = (
                f"A1:{get_column_letter(worksheet.max_column)}"
                f"{worksheet.max_row}"
            )

            table = Table(
                displayName="LeadReport",
                ref=table_ref,
            )

            table_style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )

            table.tableStyleInfo = table_style
            worksheet.add_table(table)

        for index, field in enumerate(self.FIELDNAMES, start=1):
            column_letter = get_column_letter(index)

            worksheet.column_dimensions[
                column_letter
            ].width = self.WIDTHS.get(field, 18)

        for row in range(2, worksheet.max_row + 1):
            created_at_cell = worksheet.cell(
                row=row,
                column=self.FIELDNAMES.index("created_at") + 1,
            )

            if created_at_cell.value:
                created_at_cell.number_format = (
                    "yyyy-mm-dd hh:mm"
                )

        self._add_hyperlinks(worksheet)
        self._format_scores(worksheet)
        self._format_validation_status(worksheet)

    def _add_hyperlinks(self, worksheet) -> None:
        """Turn URL fields into readable clickable hyperlinks."""

        for field in self.LINK_COLUMNS:
            column = self.FIELDNAMES.index(field) + 1

            for row in range(2, worksheet.max_row + 1):
                cell = worksheet.cell(
                    row=row,
                    column=column,
                )

                if not cell.value:
                    continue

                cell.hyperlink = cell.value

                if field == "website":
                    cell.value = "Open Website"
                elif field == "linkedin_url":
                    cell.value = "LinkedIn"
                else:
                    cell.value = "View Source"

    def _format_scores(self, worksheet) -> None:
        """Add conditional formatting to lead scores."""

        column = self.FIELDNAMES.index("lead_score") + 1
        column_letter = get_column_letter(column)

        if worksheet.max_row < 2:
            return

        score_range = (
            f"{column_letter}2:"
            f"{column_letter}{worksheet.max_row}"
        )

        worksheet.conditional_formatting.add(
            score_range,
            ColorScaleRule(
                start_type="num",
                start_value=0,
                start_color="F8696B",
                mid_type="num",
                mid_value=50,
                mid_color="FFEB84",
                end_type="num",
                end_value=100,
                end_color="63BE7B",
            ),
        )

    def _format_validation_status(self, worksheet) -> None:
        """Highlight valid and invalid lead statuses."""

        column = self.FIELDNAMES.index(
            "validation_status"
        ) + 1
        column_letter = get_column_letter(column)

        if worksheet.max_row < 2:
            return

        status_range = (
            f"{column_letter}2:"
            f"{column_letter}{worksheet.max_row}"
        )

        from openpyxl.formatting.rule import CellIsRule

        worksheet.conditional_formatting.add(
            status_range,
            CellIsRule(
                operator="equal",
                formula=['"valid"'],
                fill=PatternFill(
                    fill_type="solid",
                    fgColor="C6EFCE",
                ),
            ),
        )

        worksheet.conditional_formatting.add(
            status_range,
            CellIsRule(
                operator="equal",
                formula=['"invalid"'],
                fill=PatternFill(
                    fill_type="solid",
                    fgColor="FFC7CE",
                ),
            ),
        )